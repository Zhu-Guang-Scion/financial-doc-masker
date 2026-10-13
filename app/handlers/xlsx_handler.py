# -*- coding: utf-8 -*-
"""Excel(.xlsx) 处理器（openpyxl，单元格级处理）。

- data_only=False 加载：默认保留全部公式；
- 数值单元格与文本单元格分流：数值走金额脱敏专项（engine.mask_numeric），
  文本走统一规则引擎；
- TODO: 富文本单元格 / 批注 / 图表内文字。
"""
from __future__ import annotations

from pathlib import Path

from ..utils import MissingDependencyError, validate_id_card, validate_phone
from .text_replacement import apply_replacements


def contact_numeric_replacement(v, engine, options):
    """数值格手机号/证号独立通道（返回 (规则, 替换) 或 None）。

    数值格此前唯一出口是金额通道（mask_amount 默认 False → 默认配置下
    数值格手机号零通道覆盖直接裸奔）。本通道在
    金额通道之前独立识别：整数值规范成字符串后走 validate_phone /
    validate_id_card 校验，与文本通道同门 mask_contact、独立于 mask_amount。
    注意：18 位证号在 Excel float 下有效数字只剩 15 位、尾位已被 Excel
    破坏，validate_id_card 校验位会天然拒认——该形态本就无保真原文可救，
    实际受益面=11 位手机号（<15 位安全）。
    """
    if not getattr(options, "mask_contact", True) or isinstance(v, bool):
        return None
    if not isinstance(v, (int, float)) or not float(v).is_integer():
        return None
    s = str(int(v))
    if validate_phone(s):
        return ("phone", engine._default_replacement("phone", s))
    if validate_id_card(s):
        return ("id_card", engine._default_replacement("id_card", s))
    return None

# 表头含这些关键词的列视为“人名列”：列内非空文本强制按人名脱敏
# （序时账/台账的制单人、审核人等与姓名分列存放，上下文规则无法覆盖）
_PERSON_HEADER_KEYS = ("姓名", "制单", "审核", "经办", "复核", "审批", "出纳",
                       "会计", "财务负责人", "联系人", "收款", "付款", "借款",
                       "报销", "申请", "签字", "领款", "法人", "监事", "董事",
                       "股东", "出质", "抵押", "保证", "司库", "填报", "填表")


def _detect_person_columns(ws) -> set:
    """扫描前两行表头，返回命中人名关键词的列号集合。"""
    cols = set()
    for row in ws.iter_rows(min_row=1, max_row=min(2, ws.max_row or 1)):
        for cell in row:
            if isinstance(cell.value, str):
                if any(k in cell.value for k in _PERSON_HEADER_KEYS):
                    cols.add(cell.column)
    return cols


def process(src: Path, dst: Path, engine, options, report) -> dict:
    try:
        from openpyxl import load_workbook
        from openpyxl.cell.cell import MergedCell
    except ImportError as e:
        raise MissingDependencyError(
            f"处理 .xlsx 需要 openpyxl：pip install openpyxl。组件已安装仍报此错时，多为依赖库 DLL 加载失败（请安装 Microsoft VC++ 2015-2022 x64 运行库后重试）；内部原因：{e}"
        ) from e

    report.log("阶段 1/2：加载工作簿（保留公式）…")
    wb = load_workbook(str(src), data_only=False)
    hits = 0
    sheet_total = len(wb.worksheets)

    for si, ws in enumerate(wb.worksheets, 1):
        sheet_hits = 0
        person_cols = _detect_person_columns(ws)  # 表头驱动的人名列
        if person_cols:
            report.log(f"工作表[{ws.title}]：检测到人名列（表头驱动，强制脱敏）")
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or isinstance(cell, MergedCell):
                    continue
                force = cell.column in person_cols and cell.row > 1
                # 公式单元格默认原样保留。
                # TODO: 实现“保留公式仅脱敏显示值”模式
                #       （需引入公式计算引擎，或改写公式引用常量后重算）。
                if cell.data_type == "f":
                    continue
                v = cell.value
                if isinstance(v, bool):
                    continue
                if isinstance(v, (int, float)):
                    ct = contact_numeric_replacement(v, engine, options)
                    if ct is not None:  # 数值格手机号/证号独立通道
                        rule, repl = ct
                        cell.value = repl
                        report.record(str(int(v)), repl, rule,
                                      f"[{ws.title}]!{cell.coordinate}")
                        sheet_hits += 1
                        continue
                    new_v = engine.mask_numeric(v)  # 数值单元格：金额脱敏专项
                    if new_v is not None:
                        # TODO: 写入字符串会丢失数字格式；可结合 number_format 保留显示样式
                        cell.value = new_v
                        sheet_hits += 1
                elif isinstance(v, str):
                    spans = engine.find_spans(v, force_person=force)
                    if spans:
                        for sp in spans:
                            report.record(sp.text, sp.replacement, sp.rule,
                                          f"[{ws.title}]!{cell.coordinate}")
                        cell.value = apply_replacements(v, spans)
                        sheet_hits += len(spans)
        if sheet_hits:
            report.log(f"工作表[{ws.title}]：脱敏 {sheet_hits} 处")
        hits += sheet_hits
        report.progress(si, sheet_total)

    report.log(f"阶段 2/2：写出脱敏副本（命中 {hits} 处）…")
    wb.save(str(dst))
    # TODO: openpyxl 重写会丢弃图表/部分富文本等高级特性；
    #       如需 100% 保留格式，可仿照 xls_handler 走本机 Excel/WPS 的 COM 链路。
    return {"masked": hits, "sheets": sheet_total}
