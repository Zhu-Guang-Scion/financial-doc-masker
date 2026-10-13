# -*- coding: utf-8 -*-
"""PPT(.pptx) 处理器（python-pptx）。

覆盖：文本框、表格、组合形状（递归）。
TODO: 母版/版式/备注页文字、图表内文字、SmartArt。
"""
from __future__ import annotations

from pathlib import Path

from ..utils import MissingDependencyError
from .text_replacement import rewrite_runs


def process(src: Path, dst: Path, engine, options, report) -> dict:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as e:
        raise MissingDependencyError(
            f"处理 .pptx 需要 python-pptx：pip install python-pptx。组件已安装仍报此错时，多为依赖库 DLL 加载失败（请安装 Microsoft VC++ 2015-2022 x64 运行库后重试）；内部原因：{e}"
        ) from e

    report.log("阶段 1/2：解析 PPT 结构 …")
    prs = Presentation(str(src))
    counter = {"hits": 0}
    slides = list(prs.slides)

    def handle_paragraph(paragraph, location: str) -> None:
        text = "".join(run.text or "" for run in paragraph.runs)
        if not text.strip():
            return
        spans = engine.find_spans(text)
        if spans:
            for sp in spans:
                report.record(sp.text, sp.replacement, sp.rule, location)
            rewrite_runs(paragraph.runs, spans)
            counter["hits"] += len(spans)

    def walk_shapes(shapes, location: str) -> None:
        for idx, shape in enumerate(shapes, 1):
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:  # 组合形状递归
                    walk_shapes(shape.shapes, location)
                    continue
            except Exception as e:
                # v2.1.9：形状类型读取失败此前静默跳过——若该
                # 形状实为组合形状，其内部文本框未经脱敏且不留任何痕迹。补
                # 一行留痕（含形状名/序号，控制流语义不变：照旧不递归、按
                # 普通形状继续处理）。留痕自身再防御，绝不反向引入新异常。
                try:
                    _name = ""
                    try:
                        _name = str(getattr(shape, "name", "") or "")
                    except Exception:
                        pass
                    _where = f"形状「{_name}」" if _name else f"第{idx}个形状"
                    report.log(f"⚠ {location}{_where}类型读取失败已跳过"
                               f"（[{type(e).__name__}] {e}）：若为组合形状，"
                               f"内部文本框未经脱敏，请人工复核。")
                except Exception:
                    pass
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    handle_paragraph(para, location)
            if getattr(shape, "has_table", False) and shape.has_table:
                for ri, row in enumerate(shape.table.rows, 1):
                    for ci, cell in enumerate(row.cells, 1):
                        for para in cell.text_frame.paragraphs:
                            handle_paragraph(para, f"{location} 表格R{ri}C{ci}")

    for i, slide in enumerate(slides, 1):
        walk_shapes(slide.shapes, f"第{i}页")
        report.progress(i, len(slides))

    report.log(f"共 {len(slides)} 页，命中敏感信息 {counter['hits']} 处")
    report.log("阶段 2/2：写出脱敏副本 …")
    prs.save(str(dst))
    return {"masked": counter["hits"], "slides": len(slides)}
