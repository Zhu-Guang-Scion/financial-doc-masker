# -*- coding: utf-8 -*-
"""v2.1.7 批乙 P1-8：Excel 数值格手机号/证号独立通道（先红后绿）。

背景：数值格此前唯一出口是金额通道（mask_amount 默认 False → 默认配置
下数值格手机号零通道覆盖直接裸奔）。本通道独立于 mask_amount，与文本
通道同门 mask_contact。

例值均为合成构造（连号假号），不涉真实语料。
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

from app.handlers import xlsx_handler  # noqa: E402
from app.rules_engine import RuleEngine, MaskOptions  # noqa: E402


class _FakeReport:
    def __init__(self):
        self.records = []
        self.logs = []

    def record(self, original, replacement, rule, location):
        self.records.append((original, replacement, rule, location))

    def log(self, msg):
        self.logs.append(msg)

    def progress(self, *a):
        pass


def _build_book(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "台账"
    ws["A1"] = 13800001234      # 数值格手机号（默认口径下裸奔的形态）
    ws["A2"] = 12345            # 普通小数值：不得被本通道动
    ws["A3"] = "13800001234"    # 文本格手机号：既有行为回归锚
    ws["A4"] = 13800001234.5    # 非整数：跳过
    wb.save(path)


def _run(tmp_path, **opt_kw):
    src = tmp_path / "样例.xlsx"
    dst = tmp_path / "样例_脱敏后.xlsx"
    _build_book(src)
    kw = {"mask_amount": False, "output_dir": str(tmp_path)}
    kw.update(opt_kw)
    options = MaskOptions(**kw)
    engine = RuleEngine(options)
    report = _FakeReport()
    stats = xlsx_handler.process(src, dst, engine, options, report)
    wb = openpyxl.load_workbook(dst)
    ws = wb["台账"]
    return ws, report, stats


def test_numeric_cell_phone_masked_default_profile(tmp_path):
    """红用例：mask_amount=False（默认口径）下数值格手机号必须被遮蔽。"""
    ws, report, stats = _run(tmp_path)
    v1 = ws["A1"].value
    assert v1 != 13800001234, f"数值格手机号未被遮蔽：{v1!r}"
    assert isinstance(v1, str) and v1, "遮蔽后应为非空替换文本"
    rules = [r[2] for r in report.records]
    assert "phone" in rules, f"应产生 phone 规则记录：{report.records}"
    assert stats.get("masked", 0) >= 1


def test_numeric_cell_normal_values_untouched(tmp_path):
    ws, report, _ = _run(tmp_path)
    assert ws["A2"].value == 12345, "普通小数值不得被本通道改动"
    assert ws["A4"].value == 13800001234.5, "非整数数值应原样保留"


def test_text_cell_phone_still_masked(tmp_path):
    """回归锚：文本格手机号沿统一规则引擎照旧遮蔽。"""
    ws, report, _ = _run(tmp_path)
    assert ws["A3"].value != "13800001234", "文本格手机号遮蔽行为回退"
    assert any(r[0] == "13800001234" and r[2] == "phone" for r in report.records)


def test_numeric_phone_with_amount_on_still_masked(tmp_path):
    """mask_amount=True 时数值格手机号行为不回退（通道独立）。"""
    ws, report, _ = _run(tmp_path, mask_amount=True)
    assert ws["A1"].value != 13800001234
    assert any(r[2] == "phone" for r in report.records)


def test_contact_off_numeric_phone_passthrough(tmp_path):
    """mask_contact=False：通道关闭，数值格行为与历史一致（金额通道接管）。"""
    ws, report, _ = _run(tmp_path, mask_contact=False)
    assert not any(r[2] == "phone" and r[3].startswith("[台账]!A1")
                   for r in report.records)


def test_helper_contract():
    """helper 纯函数契约：整数手机号/证号/普通数三分支。"""
    eng = RuleEngine(MaskOptions())
    from app.handlers.xlsx_handler import contact_numeric_replacement

    class _Opt:
        mask_contact = True

    r = contact_numeric_replacement(13800001234, eng, _Opt())
    assert r and r[0] == "phone" and r[1]
    # 18 位证号构造值（合成）：校验位判定为准——合法则 id_card 通道，
    # 非法（或 float 精度损坏形态）则 None 交还金额通道/原样，两态都合法
    r2 = contact_numeric_replacement(int("110101199003079970"), eng, _Opt())
    assert r2 is None or (isinstance(r2, tuple) and r2[0] == "id_card")
    assert contact_numeric_replacement(12345, eng, _Opt()) is None
    assert contact_numeric_replacement(13800001234.5, eng, _Opt()) is None
    assert contact_numeric_replacement(True, eng, _Opt()) is None

    class _Off:
        mask_contact = False

    assert contact_numeric_replacement(13800001234, eng, _Off()) is None
